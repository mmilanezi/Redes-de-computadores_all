#!/usr/bin/env python3
"""
Gera a apresentacao do Trabalho I sobre o modelo oficial da UNISINOS.

Cada slide do modelo e uma imagem de fundo em tela cheia (logo e moldura ja
desenhados) com caixas de texto por cima. O gerador reaproveita essas imagens
de fundo e monta o conteudo com formas, tabelas e graficos NATIVOS do
PowerPoint, lendo os numeros de resultados/metricas.json.

Uso:  python scripts/gerar_apresentacao.py
Saida: Apresentacao_Trabalho1_Matheus_Milanezi.pptx
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
MODELO = RAIZ / "Modelo-de-apresentacao-Unisinos.pptx"
# os dados ficam no repositorio entregue ao professor; este gerador e o
# modelo da UNISINOS ficam fora dele
DADOS = RAIZ / "repo" / "resultados" / "metricas.json"
SAIDA = RAIZ / "Apresentacao_Trabalho1_Matheus_Milanezi.pptx"

# ---------------------------------------------------------------------------
# Identidade visual: azul do modelo + as cores de protocolo dos graficos
# ---------------------------------------------------------------------------
AZUL = RGBColor(0x25, 0x3A, 0x7E)        # cor institucional do modelo
AZUL_ESCURO = RGBColor(0x1A, 0x2A, 0x5E)
TEXTO_2 = RGBColor(0x4A, 0x55, 0x78)     # texto secundario, azul acinzentado
MUDO = RGBColor(0x6B, 0x72, 0x87)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CARTAO = RGBColor(0xEE, 0xF1, 0xF8)      # tinta clara do azul do modelo
CARTAO_BORDA = RGBColor(0xD5, 0xDB, 0xEA)
DESTAQUE_OSPF = RGBColor(0xFB, 0xED, 0xE6)
LINHA = RGBColor(0x9A, 0xA3, 0xB8)
GRADE = RGBColor(0xDD, 0xE1, 0xEA)
COR = {"rip": RGBColor(0x2A, 0x78, 0xD6),
       "ospf": RGBColor(0xEB, 0x68, 0x34),
       "bgp": RGBColor(0x1B, 0xAF, 0x7A)}
ROTULO = {"rip": "RIPv2", "ospf": "OSPFv2", "bgp": "BGP-4"}

# Imagens de fundo do modelo (ppt/media/imageN.png)
FUNDO_CAPA = "ppt/media/image1.png"         # capa branca com o U azul
FUNDO_CONTEUDO = "ppt/media/image3.png"     # branco com logo e moldura
FUNDO_DESTAQUE = "ppt/media/image4.png"     # azul com logo e moldura
FUNDO_FIM = "ppt/media/image6.png"          # encerramento branco

# Grade dos slides de conteudo (polegadas; slide de 13,33 x 7,5)
X0 = 1.88           # alinhado com o texto do modelo, a direita do logo
LARG = 10.87        # ate a moldura da direita
Y_CONTEUDO = 2.0


def limpar_fundo_conteudo(dados: bytes) -> bytes:
    """
    A imagem de fundo branca do modelo traz um residuo: o hifen de 'laboris-'
    do texto de exemplo ficou gravado nela, em (1647-1655, 533-534) px. Onde o
    slide nao tem nada por cima, ele aparece como um traco solto. Pinta a
    regiao com a cor do fundo ao redor.
    """
    from PIL import Image
    im = Image.open(io.BytesIO(dados)).convert("RGBA")
    cor_fundo = im.getpixel((1630, 533))
    for y in range(528, 540):
        for x in range(1642, 1661):
            im.putpixel((x, y), cor_fundo)
    saida = io.BytesIO()
    im.save(saida, format="PNG")
    return saida.getvalue()


# ===========================================================================
# Primitivas
# ===========================================================================

def novo_slide(prs, fundo: str, imagens: dict):
    layout = next(l for l in prs.slide_layouts if l.name == "Em branco")
    s = prs.slides.add_slide(layout)
    s.shapes.add_picture(io.BytesIO(imagens[fundo]), 0, 0,
                         prs.slide_width, prs.slide_height)
    return s


def texto(slide, x, y, w, h, paragrafos, *, tamanho=16, cor=AZUL,
          negrito=False, alinhamento=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP,
          fonte=None, espaco_depois=0):
    """
    paragrafos: str, ou lista de paragrafos; cada paragrafo e str ou lista de
    trechos (texto, {opcoes}) com opcoes negrito/cor/tamanho/fonte.
    """
    caixa = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caixa.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancora
    _preencher(tf, paragrafos, tamanho, cor, negrito, alinhamento, fonte,
               espaco_depois)
    return caixa


def _preencher(tf, paragrafos, tamanho, cor, negrito, alinhamento, fonte,
               espaco_depois):
    if isinstance(paragrafos, str):
        paragrafos = [paragrafos]
    for i, par in enumerate(paragrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhamento
        if espaco_depois and i < len(paragrafos) - 1:
            p.space_after = Pt(espaco_depois)
        trechos = [(par, {})] if isinstance(par, str) else par
        for conteudo, op in trechos:
            r = p.add_run()
            r.text = conteudo
            f = r.font
            f.size = Pt(op.get("tamanho", tamanho))
            f.bold = op.get("negrito", negrito)
            f.color.rgb = op.get("cor", cor)
            if op.get("fonte", fonte):
                f.name = op.get("fonte", fonte)


def titulo(slide, principal: str, apoio: str, *, claro=False):
    """Par titulo em negrito + linha regular, como no modelo."""
    cor = BRANCO if claro else AZUL
    texto(slide, X0, 0.62, LARG, 0.7, principal, tamanho=34, negrito=True,
          cor=cor)
    texto(slide, X0, 1.27, LARG, 0.45, apoio, tamanho=20, cor=cor)


def rodape(slide, conteudo: str, *, claro=False):
    texto(slide, X0, 6.55, LARG, 0.3, conteudo, tamanho=12,
          cor=RGBColor(0xC9, 0xD1, 0xE8) if claro else MUDO)


def cartao(slide, x, y, w, h, cabecalho, corpo, *, tam_cab=19, tam_corpo=15,
           fonte_cab=None):
    forma = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                   Inches(y), Inches(w), Inches(h))
    forma.adjustments[0] = 0.07
    forma.fill.solid()
    forma.fill.fore_color.rgb = CARTAO
    forma.line.fill.background()
    forma.shadow.inherit = False
    tf = forma.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    _preencher(tf, [[(cabecalho, {"negrito": True, "tamanho": tam_cab,
                                  "fonte": fonte_cab})],
                    [(corpo, {"tamanho": tam_corpo, "cor": TEXTO_2})]],
               tam_corpo, AZUL, False, PP_ALIGN.LEFT, None, 6)
    return forma


def tabela(slide, x, y, larguras, linhas, *, tamanho=16, altura_linha=0.44,
           destaque=None, cab_fundo=AZUL, cab_texto=BRANCO, corpo_fundo=None,
           corpo_texto=AZUL):
    """linhas[0] e o cabecalho. destaque: {indice_linha: cor_de_fundo}."""
    n_l, n_c = len(linhas), len(larguras)
    gt = slide.shapes.add_table(n_l, n_c, Inches(x), Inches(y),
                                Inches(sum(larguras)),
                                Inches(altura_linha * n_l))
    t = gt.table
    t.first_row = True
    t.horz_banding = False
    for j, lw in enumerate(larguras):
        t.columns[j].width = Inches(lw)
    for i in range(n_l):
        t.rows[i].height = Inches(altura_linha)
        for j in range(n_c):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.12)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = cab_fundo
            elif destaque and i in destaque:
                c.fill.fore_color.rgb = destaque[i]
            else:
                c.fill.fore_color.rgb = corpo_fundo or (
                    BRANCO if i % 2 else RGBColor(0xF6, 0xF7, 0xFB))
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(linhas[i][j])
            r.font.size = Pt(tamanho)
            r.font.bold = (i == 0) or (j == 0)
            r.font.color.rgb = cab_texto if i == 0 else corpo_texto
    return gt


def linha(slide, x1, y1, x2, y2, cor, espessura, *, tracejado=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = cor
    c.line.width = Pt(espessura)
    if tracejado:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def retangulo(slide, x, y, w, h, cor, *, arredondado=False, borda=None,
              raio=0.12):
    tipo = MSO_SHAPE.ROUNDED_RECTANGLE if arredondado else MSO_SHAPE.RECTANGLE
    f = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    if arredondado:
        f.adjustments[0] = raio
    f.fill.solid()
    f.fill.fore_color.rgb = cor
    if borda:
        f.line.color.rgb = borda
        f.line.width = Pt(1)
    else:
        f.line.fill.background()
    f.shadow.inherit = False
    return f


def notas(slide, conteudo: str):
    slide.notes_slide.notes_text_frame.text = conteudo


def fmt(v: float, casas: int = 1) -> str:
    return f"{v:.{casas}f}".replace(".", ",")


# ===========================================================================
# Slides
# ===========================================================================

def s_capa(prs, img):
    s = novo_slide(prs, FUNDO_CAPA, img)
    x, w = 5.76, 6.95
    texto(s, x, 2.05, w, 0.4, "Trabalho I", tamanho=20)
    texto(s, x, 2.45, w, 1.45, "RIP, OSPF e BGP sobre a mesma rede",
          tamanho=42, negrito=True)
    texto(s, x, 3.95, w, 0.6, "Redes de Computadores", tamanho=32)
    texto(s, x, 4.7, w, 0.9,
          ["Matheus Milanezi", "Prof. Cristiano Bonato Both · 28/09/2026"],
          tamanho=18, espaco_depois=2)
    notas(s, "Abertura, uns 15 segundos. Um laboratório com cinco roteadores "
             "FRRouting em três Sistemas Autônomos. Configurei RIP, OSPF e BGP "
             "sobre exatamente a mesma topologia, um de cada vez, e medi como "
             "cada um se comporta. Azul é RIP, laranja é OSPF, verde é BGP em "
             "todos os gráficos.")


def s_construido(prs, img):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Um laboratório de roteamento reproduzível",
           "O que foi construído")
    w, h, g = 5.3, 2.0, 0.27
    itens = [
        ("FRRouting 8.4.1",
         "O fork mantido do Quagga citado no enunciado, rodando em containers "
         "Docker."),
        ("5 roteadores, 3 Sistemas Autônomos",
         "Um anel com uma rede de acesso e um host em cada roteador."),
        ("3 protocolos, um por vez",
         "O container sobe um único daemon de roteamento. Não há como dois "
         "coexistirem."),
        ("6 experimentos automatizados",
         "Convergência, tabela, seleção de rota, tráfego de controle, "
         "reconvergência e latência."),
    ]
    for i, (cab, corpo) in enumerate(itens):
        cartao(s, X0 + (i % 2) * (w + g), Y_CONTEUDO + (i // 2) * (h + g),
               w, h, cab, corpo, tam_cab=22, tam_corpo=17)
    notas(s, "Uns 40 segundos. A plataforma é o FRRouting, sucessor mantido do "
             "Quagga, que o enunciado cita. Cada roteador é um container. O "
             "enunciado pede que os protocolos não sejam simultâneos: isso é "
             "garantido na origem, o container só liga um daemon de roteamento "
             "por vez. E a coleta de métricas é automatizada, os mesmos seis "
             "experimentos para cada protocolo.")


def s_plataforma(prs, img):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Como o laboratório sobe",
           "Instalação e configuração da plataforma")
    etapas = [
        ("gerar_topologia.py",
         "Uma única tabela de topologia gera o docker-compose e as 15 "
         "configurações FRR."),
        ("entrypoint-router.sh",
         "Renomeia as interfaces pelo IP, remove a rota default e liga um "
         "único daemon."),
        ("lab.sh up ospf",
         "Sobe 10 containers e prepara o host Docker para o multicast de RIP e "
         "OSPF."),
    ]
    w, seta, g = 3.15, 0.42, 0.14
    x = X0
    for i, (cab, corpo) in enumerate(etapas):
        cartao(s, x, Y_CONTEUDO + 0.1, w, 2.3, cab, corpo, tam_cab=17,
               tam_corpo=16, fonte_cab="Consolas")
        x += w
        if i < len(etapas) - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                   Inches(x + g), Inches(Y_CONTEUDO + 1.05),
                                   Inches(seta), Inches(0.4))
            a.fill.solid()
            a.fill.fore_color.rgb = AZUL
            a.line.fill.background()
            a.shadow.inherit = False
            x += seta + 2 * g
    texto(s, X0, 4.85, LARG, 1.1,
          [[("O Docker nomeia as interfaces em ordem não garantida. ", {}),
            ("Renomeá-las pelo IP deixa as configurações legíveis e estáveis: ",
             {}),
            ("lan, to-r2, to-r5.", {"negrito": True})]],
          tamanho=18, cor=TEXTO_2)
    notas(s, "Uns 40 segundos. Três peças. Um gerador em Python é a fonte única "
             "da topologia: dele saem o docker-compose e as quinze "
             "configurações, então endereçamento e vizinhanças nunca divergem. "
             "O entrypoint de cada roteador renomeia as interfaces pelo IP, "
             "porque o Docker não garante a ordem eth0, eth1; remove a rota "
             "default; e liga só o daemon do protocolo escolhido. E um script "
             "que sobe tudo com um comando.")


def s_topologia(prs, img):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Um anel com uma armadilha proposital",
           "Topologia física e lógica")

    # regioes dos ASes (primeiro, para ficarem atras)
    retangulo(s, 3.6, 1.98, 3.35, 2.27, CARTAO, arredondado=True,
              borda=CARTAO_BORDA, raio=0.08)
    retangulo(s, 2.5, 5.35, 4.65, 1.1, CARTAO, arredondado=True,
              borda=CARTAO_BORDA, raio=0.15)
    # o AS 65003 se estende para a esquerda do R5: embaixo dele, o rotulo
    # seria cortado pelo enlace R5-R4
    retangulo(s, 1.0, 3.3, 2.4, 0.96, CARTAO, arredondado=True,
              borda=CARTAO_BORDA, raio=0.15)

    centros = {"r1": (4.40, 2.50), "r2": (6.16, 3.78), "r3": (5.49, 5.85),
               "r4": (3.31, 5.85), "r5": (2.64, 3.78)}
    anel = [("r1", "r2"), ("r2", "r3"), ("r3", "r4"), ("r4", "r5"),
            ("r5", "r1")]
    for a, b in anel:
        (x1, y1), (x2, y2) = centros[a], centros[b]
        caro = {a, b} == {"r2", "r3"}
        linha(s, x1, y1, x2, y2, COR["ospf"] if caro else LINHA,
              5.5 if caro else 3)

    bw, bh = 1.25, 0.62
    for r, (cx, cy) in centros.items():
        caixa = retangulo(s, cx - bw / 2, cy - bh / 2, bw, bh, AZUL,
                          arredondado=True, raio=0.18)
        tf = caixa.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _preencher(tf, [[(r.upper(), {"negrito": True, "tamanho": 18})],
                        [(f"10.{r[1]}.0.0/24", {"tamanho": 11,
                                                "cor": RGBColor(0xC9, 0xD1,
                                                                0xE8)})]],
                   18, BRANCO, False, PP_ALIGN.CENTER, None, 0)

    texto(s, 5.4, 2.06, 1.45, 0.3, "AS 65001", tamanho=13, negrito=True,
          cor=TEXTO_2, alinhamento=PP_ALIGN.RIGHT)
    texto(s, 6.2, 5.75, 0.9, 0.3, "AS 65002", tamanho=13, negrito=True,
          cor=TEXTO_2)
    texto(s, 1.1, 3.4, 0.85, 0.8, ["AS", "65003"], tamanho=13, negrito=True,
          cor=TEXTO_2, ancora=MSO_ANCHOR.MIDDLE)
    texto(s, 6.12, 4.5, 1.35, 0.6,
          [[("custo OSPF 200", {"negrito": True})],
           [("demais: 10", {"tamanho": 12, "cor": TEXTO_2})]],
          tamanho=14)

    blocos = [
        ("Dois caminhos entre quaisquer pontos",
         "A queda de qualquer enlace mantém a rede conexa, o que permite medir "
         "reconvergência."),
        ("Fronteiras eBGP e sessões iBGP",
         "eBGP em R2–R3, R4–R5 e R5–R1. iBGP em R1–R2 e R3–R4."),
        ("A armadilha",
         "R2–R3 é o caminho mais curto em saltos entre os dois ASes, mas o mais "
         "caro para o OSPF. RIP e OSPF vão discordar."),
    ]
    y = 2.05
    for cab, corpo in blocos:
        texto(s, 7.65, y, 5.1, 1.3,
              [[(cab, {"negrito": True, "tamanho": 19})],
               [(corpo, {"tamanho": 16, "cor": TEXTO_2})]],
              espaco_depois=4)
        y += 1.45
    rodape(s, "Enlaces /29, LANs /24, redes internas sem NAT e sem rota default")
    notas(s, "Uns 50 segundos. Os cinco roteadores formam um anel: entre "
             "quaisquer dois pontos há dois caminhos. R1 e R2 são o AS 65001, "
             "R3 e R4 o 65002, R5 sozinho o 65003. A decisão de projeto "
             "importante está em laranja: o enlace R2–R3 é o caminho mais curto "
             "em saltos entre o AS 65001 e o 65002, mas tem custo OSPF 200 "
             "contra 10 dos demais. Fiz isso para que RIP e OSPF escolham "
             "caminhos diferentes. As redes são internas, sem NAT e sem rota "
             "default: todo o encaminhamento depende só do protocolo.")


def s_selecao(prs, img, d):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Mesmo destino, caminhos diferentes",
           "Seleção de rota de R1 até a LAN de R3 (10.3.0.0/24)")
    e3 = {p: d[p]["e3_selecao_rota"] for p in d}
    linhas = [
        ["Protocolo", "Caminho", "Métrica", "Critério"],
        ["RIPv2", "R1 → R2 → R3", f"{e3['rip']['metrica']} saltos",
         "menos saltos"],
        ["OSPFv2", "R1 → R5 → R4 → R3", f"custo {e3['ospf']['metrica']}",
         "menor custo"],
        ["BGP-4", "R1 → R2 → R3", "AS_PATH 65002", "AS_PATH mais curto"],
    ]
    tabela(s, X0, Y_CONTEUDO, [1.9, 3.5, 2.4, 3.07], linhas, tamanho=19,
           altura_linha=0.52, destaque={2: DESTAQUE_OSPF})
    cartao(s, X0, 4.4, 5.3, 1.85, "O OSPF dá a volta",
           "Pelo R2 o custo seria 10 + 200 + 10 = 220. Pela volta, 40. Um "
           "roteador a mais para evitar o enlace caro.", tam_cab=20,
           tam_corpo=16)
    cartao(s, X0 + 5.57, 4.4, 5.3, 1.85,
           "O BGP coincide com o RIP só no resultado",
           "Ele compara AS_PATH 65002 contra 65003 65002 e não enxerga os "
           "roteadores dentro de cada AS.", tam_cab=20, tam_corpo=16)
    rodape(s, "Fonte: show ip route, show ip bgp e traceroute h1 → h3")
    notas(s, "Resultado central, uns 60 segundos. Mesmo destino, a LAN do R3, "
             "saindo do R1. O RIP vai pelo R2: conta saltos e não sabe que "
             "aquele enlace é caro. O OSPF dá a volta pelo R5 e R4: pelo R2 o "
             "custo seria 220, pela volta é 40. O BGP também vai pelo R2, por "
             "outro motivo: tem dois caminhos, AS_PATH 65002 direto e 65003 "
             "65002 pelo AS do R5, e escolhe o mais curto. Se o AS 65002 "
             "tivesse vinte roteadores internos, o BGP continuaria vendo um AS "
             "só. Coincidência de resultado, não de critério.")


def s_convfria(prs, img, d):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "O OSPF é o mais lento para começar",
           "Convergência a frio: até todos os roteadores conhecerem as 5 LANs")

    dados = XyChartData()
    tmax = 0
    for p in ("rip", "ospf", "bgp"):
        serie = dados.add_series(ROTULO[p])
        for t, n in d[p]["e1_convergencia_fria"]["curva"]:
            serie.add_data_point(t, n)
            tmax = max(tmax, t)
    completa = dados.add_series("convergência completa")
    completa.add_data_point(0, 25)
    completa.add_data_point(60, 25)

    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
                            Inches(X0 - 0.1), Inches(Y_CONTEUDO - 0.1),
                            Inches(7.1), Inches(4.55), dados)
    ch = gf.chart
    ch.font.size = Pt(12)
    ch.font.color.rgb = TEXTO_2
    cores = [COR["rip"], COR["ospf"], COR["bgp"], LINHA]
    for serie, cor in zip(ch.plots[0].series, cores):
        serie.smooth = False
        serie.marker.style = XL_MARKER_STYLE.NONE
        serie.format.line.color.rgb = cor
        serie.format.line.width = Pt(2.75)
    ch.plots[0].series[3].format.line.width = Pt(1.25)
    ch.plots[0].series[3].format.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    eixo_x, eixo_y = ch.category_axis, ch.value_axis
    eixo_x.minimum_scale, eixo_x.maximum_scale, eixo_x.major_unit = 0, 60, 10
    eixo_y.minimum_scale, eixo_y.maximum_scale, eixo_y.major_unit = 0, 28, 5
    eixo_x.has_major_gridlines = False
    eixo_y.has_major_gridlines = True
    eixo_y.major_gridlines.format.line.color.rgb = GRADE
    for eixo, nome in ((eixo_x, "segundos desde a subida dos daemons"),
                       (eixo_y, "rotas de acesso conhecidas (máx. 25)")):
        eixo.format.line.color.rgb = GRADE
        eixo.tick_labels.font.size = Pt(12)
        eixo.tick_labels.font.color.rgb = MUDO
        eixo.has_title = True
        eixo.axis_title.text_frame.text = nome
        f = eixo.axis_title.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold, f.color.rgb = Pt(12), False, TEXTO_2
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)

    e1 = {p: d[p]["e1_convergencia_fria"]["segundos"] for p in d}
    blocos = [
        (f"BGP · {e1['bgp']:.0f} s",
         "A sessão TCP sobe e os UPDATEs saem em seguida."),
        (f"RIP · {e1['rip']:.0f} s",
         "Depende do próximo update de 30 s. Em outra subida, 68 s."),
        (f"OSPF · {e1['ospf']:.0f} s",
         "40 s em Waiting antes de eleger o DR, porque os enlaces são redes "
         "broadcast."),
    ]
    y = 2.05
    for cab, corpo in blocos:
        texto(s, 9.3, y, 3.45, 1.3,
              [[(cab, {"negrito": True, "tamanho": 20})],
               [(corpo, {"tamanho": 16, "cor": TEXTO_2})]], espaco_depois=4)
        y += 1.45
    rodape(s, "Fonte: show ip route json nos 5 roteadores, a cada 1 s")
    notas(s, "Uns 45 segundos. Tempo desde a subida dos daemons até todos "
             "conhecerem todas as LANs. O BGP é o mais rápido, 4 segundos. O "
             "RIP depende do timer de 30 segundos; numa subida deu 29, em outra "
             "68, então esse número varia muito. O surpreendente é o OSPF: a "
             "linha laranja fica parada por mais de 40 segundos. Os enlaces do "
             "Docker são redes broadcast, e nelas o OSPF espera o wait timer, "
             "40 segundos, antes de eleger o DR. Enlaces ponto a ponto "
             "eliminariam essa espera.")


def s_controle(prs, img, d):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "O medido bate com os timers",
           "Tráfego de controle em regime permanente: 90 s, soma dos 5 "
           "roteadores")
    e4 = {p: d[p]["e4_trafego_controle"] for p in d}

    def bpp(p):
        return round(e4[p]["total_bytes"] / max(1, e4[p]["total_pacotes"]))

    def bits(p):
        # arredonda 0,5 para cima (o round() do Python arredonda para o par,
        # e 138,5 viraria 138 aqui e 139 no README)
        return int(e4[p]["bits_por_s"] + 0.5)

    linhas = [
        ["Protocolo", "pacotes/s", "bit/s", "bytes/pacote", "Previsto"],
        ["RIPv2", fmt(e4["rip"]["pacotes_por_s"], 3),
         bits("rip"), bpp("rip"),
         "10 interfaces × 1 update/30 s = 0,333"],
        ["OSPFv2", fmt(e4["ospf"]["pacotes_por_s"], 3),
         bits("ospf"), bpp("ospf"),
         "10 interfaces × 1 hello/10 s = 1,000"],
        ["BGP-4", fmt(e4["bgp"]["pacotes_por_s"], 3),
         bits("bgp"), bpp("bgp"),
         "10 pontas × 1 keepalive/60 s + ACKs"],
    ]
    tabela(s, X0, Y_CONTEUDO, [1.75, 1.45, 1.2, 1.75, 4.72], linhas,
           tamanho=17, altura_linha=0.48)
    w, g = 3.45, 0.26
    cartoes = [
        ("OSPF: muitos pacotes pequenos",
         "Hellos só confirmam o vizinho. O tamanho não depende da rede."),
        ("RIP: a tabela inteira a cada 30 s",
         "Já gasta 84% dos bits do OSPF. Com mais rotas, passa à frente."),
        ("BGP: o mais barato",
         "Sem mudanças, só keepalives. Updates são incrementais."),
    ]
    for i, (cab, corpo) in enumerate(cartoes):
        cartao(s, X0 + i * (w + g), 4.35, w, 1.95, cab, corpo, tam_cab=18,
               tam_corpo=15)
    rodape(s, "Fonte: tcpdump -Q out nos 5 roteadores; cada pacote contado uma "
              "vez, no emissor")
    notas(s, "Uns 50 segundos. Capturei 90 segundos de tráfego de controle em "
             "cada roteador, só na saída, para não contar cada pacote duas "
             "vezes. Os números batem com a teoria: o OSPF manda um hello a "
             "cada 10 segundos em 10 interfaces, exatamente 1 pacote por "
             "segundo; o RIP, um update a cada 30, dá 0,333. Isso valida a "
             "medição. Sobre escalabilidade: o OSPF manda mais pacotes, mas "
             "pequenos; o RIP manda poucos, mas cada um leva a tabela inteira, "
             "então o custo dele cresce com a rede. Aqui ele já gasta 84 por "
             "cento do que o OSPF gasta. O BGP em regime só troca keepalives.")


def s_tabelas(prs, img, d):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "O OSPF guarda o mapa, o BGP guarda alternativas",
           "Tamanho da tabela de roteamento")
    e2 = {p: d[p]["e2_tabelas"] for p in d}
    base_bgp = [v["base_protocolo"] for v in e2["bgp"]["por_roteador"].values()]
    linhas = [
        ["Protocolo", "FIB (instaladas)", "RIB", "Base do protocolo"],
        ["RIPv2", f"{e2['rip']['media_fib']:.0f}",
         f"{e2['rip']['media_rib']:.0f}",
         f"{e2['rip']['media_base_protocolo']:.0f} rotas"],
        ["OSPFv2", f"{e2['ospf']['media_fib']:.0f}",
         f"{e2['ospf']['media_rib']:.0f}",
         f"{e2['ospf']['media_base_protocolo']:.0f} LSAs"],
        ["BGP-4", f"{e2['bgp']['media_fib']:.0f}",
         f"{e2['bgp']['media_rib']:.0f}",
         f"{min(base_bgp)} a {max(base_bgp)} caminhos"],
    ]
    tabela(s, X0, Y_CONTEUDO, [2.2, 2.6, 2.0, 4.07], linhas, tamanho=19,
           altura_linha=0.5)
    cartao(s, X0, 4.35, 5.3, 1.95, "LSDB: a topologia",
           "5 router-LSAs e 5 network-LSAs, idênticos nos 5 roteadores. Cada "
           "um calcula as rotas a partir do mesmo mapa.", tam_cab=20,
           tam_corpo=16)
    cartao(s, X0 + 5.57, 4.35, 5.3, 1.95, "BGP: caminhos de reserva",
           f"O R5, com duas sessões eBGP, guarda {max(base_bgp)} caminhos, "
           "2 por prefixo remoto. É isso que o deixa trocar de rota na hora.",
           tam_cab=20, tam_corpo=16)
    rodape(s, "A RIB do OSPF tem 4 entradas para redes conectadas que nunca são "
              "instaladas")
    notas(s, "Uns 40 segundos. Separei o que o roteador usa para encaminhar, a "
             "FIB, do que o protocolo guarda por dentro. RIP e OSPF instalam os "
             "mesmos 15 prefixos; o BGP 12, porque cada roteador só anuncia a "
             "LAN e o loopback, não os enlaces entre ASes. Por dentro: o OSPF "
             "guarda o mapa da rede, 10 LSAs iguais em todos os roteadores. O "
             "BGP guarda caminhos alternativos: o R5, com duas sessões eBGP, "
             "conhece dois caminhos para cada prefixo. Isso aparece nos "
             "próximos slides. Na primeira versão eu contei a RIB, e o OSPF "
             "parecia ter 19 rotas; eram 4 entradas nunca instaladas.")


def s_reconv_tempo(prs, img, d):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Quando só um lado percebe a falha",
           "Ping h1 → h5 a 50 pacotes/s; o enlace R1–R5 cai só do lado do R1")

    t_min, t_max = -15, 355
    gx0, gx1 = 3.35, 12.55

    def px(t):
        return gx0 + (t - t_min) / (t_max - t_min) * (gx1 - gx0)

    y_topo, passo, alt = 2.25, 1.0, 0.5
    for t in range(0, 351, 50):
        linha(s, px(t), y_topo - 0.1, px(t), y_topo + 3 * passo - 0.25, GRADE,
              0.75)
        texto(s, px(t) - 0.4, y_topo + 3 * passo - 0.15, 0.8, 0.3, str(t),
              tamanho=12, cor=MUDO, alinhamento=PP_ALIGN.CENTER)
    texto(s, gx0, y_topo + 3 * passo + 0.15, gx1 - gx0, 0.3,
          "segundos em relação à queda do enlace", tamanho=12, cor=TEXTO_2,
          alinhamento=PP_ALIGN.CENTER)

    for i, p in enumerate(("rip", "ospf", "bgp")):
        rodada = d[p]["e5_reconvergencia"]["rodadas"][0]
        tempos = [t for t, _s in rodada["linha_do_tempo"]]
        trechos, ini = [], tempos[0]
        for a, b in zip(tempos, tempos[1:]):
            if b - a > 0.2:
                trechos.append((ini, a))
                ini = b
        trechos.append((ini, tempos[-1]))
        yc = y_topo + i * passo
        texto(s, X0, yc, 1.3, alt, ROTULO[p], tamanho=17, negrito=True,
              ancora=MSO_ANCHOR.MIDDLE)
        for a, b in trechos:
            a, b = max(a, t_min), min(b, t_max)
            retangulo(s, px(a), yc, max(0.02, px(b) - px(a)), alt, COR[p])
        volta = rodada.get("janela_enlace_caido_s")
        if volta:
            linha(s, px(volta), yc - 0.08, px(volta), yc + alt + 0.08, AZUL_ESCURO,
                  2.5)
            texto(s, px(volta) + 0.06, yc - 0.3, 1.3, 0.25, "enlace volta",
                  tamanho=11, cor=TEXTO_2)

    linha(s, px(0), y_topo - 0.25, px(0), y_topo + 3 * passo - 0.3, AZUL_ESCURO,
          1.5, tracejado=True)
    texto(s, px(0) + 0.08, y_topo - 0.42, 1.2, 0.25, "enlace cai", tamanho=11,
          cor=TEXTO_2)

    uni = {p: d[p]["e5_reconvergencia"]["media_s"] for p in d}
    texto(s, X0, 5.95, LARG, 0.5,
          [[("OSPF: nenhum pacote perdido. ", {"negrito": True}),
            (f"RIP e BGP: cerca de {int((uni['rip'] + uni['bgp']) / 2)} s "
             "parados, até vencer o timer de 180 s.", {})]], tamanho=17)
    rodape(s, "Rodada 1 de 3 em cada protocolo; o teste do OSPF usou uma janela "
              "de 58 s")
    notas(s, "Resultado mais interessante, uns 60 segundos. Derrubei o enlace "
             "R1–R5 só do lado do R1. A interface do R5 continua ativa: ele não "
             "sabe que o enlace morreu e continua mandando as respostas por "
             "ali. Como o R5 descobre? O OSPF não precisa: o R1 inunda um LSA "
             "novo sem esse enlace, e o SPF só usa um enlace se os dois lados o "
             "anunciarem. Zero pacotes perdidos. O RIP só descobre quando a "
             "rota expira, 180 segundos sem updates. O BGP, quando o hold timer "
             "vence, também 180 segundos. Por isso os dois ficam quase três "
             "minutos parados.")


def s_reconv_cenarios(prs, img, d):
    s = novo_slide(prs, FUNDO_DESTAQUE, img)
    titulo(s, "A detecção pesou mais que o protocolo",
           "Reconvergência: média de 3 rodadas por cenário", claro=True)
    uni = {p: d[p]["e5_reconvergencia"]["media_s"] for p in d}
    bil = {p: d[p]["e5b_reconvergencia_bilateral"]["media_s"] for p in d}

    def dur(v):
        return f"{v * 1000:.0f} ms" if v < 1 else f"{fmt(v)} s"

    linhas = [
        ["Protocolo", "Um lado só", "Cabo cortado", "Como o R5 descobre"],
        ["RIPv2", dur(uni["rip"]), dur(bil["rip"]), "timeout da rota, 180 s"],
        ["OSPFv2", dur(uni["ospf"]), dur(bil["ospf"]),
         "não precisa: LSA do R1 + SPF"],
        ["BGP-4", dur(uni["bgp"]), dur(bil["bgp"]), "hold timer, 180 s"],
    ]
    tabela(s, X0, Y_CONTEUDO, [2.0, 2.2, 2.3, 4.37], linhas, tamanho=19,
           altura_linha=0.52, cab_fundo=CARTAO, cab_texto=AZUL,
           corpo_fundo=BRANCO)
    fator_rip = uni["rip"] / bil["rip"]
    fator_bgp = uni["bgp"] / bil["bgp"]
    blocos = [
        (f"{fator_rip:.0f}× no RIP, {fator_bgp:.0f}× no BGP",
         "O mesmo protocolo, só mudando quem enxerga a queda. Com o cabo "
         "cortado, o BGP usa o caminho de reserva na hora."),
        ("Só o OSPF é imune",
         "Ele confirma cada enlace pelos dois lados. Nas redes reais, o BGP "
         "resolve isso com BFD."),
    ]
    for i, (cab, corpo) in enumerate(blocos):
        texto(s, X0 + i * 5.57, 4.45, 5.3, 1.8,
              [[(cab, {"negrito": True, "tamanho": 22})],
               [(corpo, {"tamanho": 17})]], cor=BRANCO, espaco_depois=6)
    rodape(s, "Falha unilateral: janela de 270 s; bilateral: 60 s", claro=True)
    notas(s, "Uns 50 segundos. Repeti tudo com o enlace cortado nas duas "
             "pontas, como um cabo rompido. O RIP cai de 175 para 12 segundos: "
             "os dois lados mandam updates disparados, mas o caminho "
             "alternativo só chega no próximo update periódico. O BGP cai de "
             "175 segundos para 125 milissegundos: as sessões caem na hora e o "
             "R5 já tinha o caminho de reserva. O OSPF é igual nos dois casos. "
             "A forma de detectar a falha pesou mais que o protocolo. Em rede "
             "real, o BGP usa BFD para isso; seria o próximo experimento.")


def s_comparativo(prs, img):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "Cada protocolo no seu lugar", "Comparação e adequação a cenários")
    linhas = [
        ["", "RIPv2", "OSPFv2", "BGP-4"],
        ["Escolhe por", "menos saltos", "menor custo", "AS_PATH e política"],
        ["Falha", "lento: timers de 30 e 180 s", "o mais rápido",
         "rápido se detecta; 180 s se não"],
        ["Controle", "tabela inteira a cada 30 s", "hellos a cada 10 s",
         "só keepalives"],
        ["Configuração", "9 linhas", "12 linhas, custos",
         "15 a 16 linhas, vizinhos"],
        ["Serve para", "redes pequenas e estáveis", "dentro de uma organização",
         "entre organizações"],
    ]
    tabela(s, X0, Y_CONTEUDO - 0.05, [2.0, 2.95, 2.95, 2.97], linhas,
           tamanho=16, altura_linha=0.47)
    texto(s, X0, 5.0, LARG, 1.2,
          [[("RIP e OSPF disputam o mesmo lugar, e o OSPF venceu onde importa. ",
             {}),
            ("O BGP não compete com eles: ", {"negrito": True}),
            ("numa rede real ele roda junto com um IGP, OSPF dentro do AS e BGP "
             "nas fronteiras.", {})]], tamanho=18, cor=TEXTO_2)
    rodape(s, "Configuração: linhas específicas do protocolo em cada roteador")
    notas(s, "Uns 40 segundos. RIP escolhe por saltos, OSPF por custo, BGP por "
             "AS_PATH e política. Em falha, o OSPF foi o melhor em todos os "
             "cenários. Em controle, o BGP é o mais barato. Em configuração, o "
             "RIP é o mais simples, 9 linhas por roteador, contra 12 do OSPF e "
             "15 a 16 do BGP, que exige conhecer cada vizinho e seu AS. RIP e "
             "OSPF competem pelo roteamento interno, e o OSPF venceu. O BGP "
             "não compete: numa rede real ele roda junto com o OSPF.")


def s_licoes(prs, img):
    s = novo_slide(prs, FUNDO_CONTEUDO, img)
    titulo(s, "O que deu errado, e o que ensinou", "Método")
    itens = [
        ("RIP e OSPF mudos no Docker",
         "O netfilter da bridge descartava o multicast. O BGP funcionava, "
         "porque usa TCP."),
        ("0,056 s nos três protocolos",
         "Valor idêntico denunciou erro de medida. Passei a buscar a maior "
         "lacuna entre respostas."),
        ("RIP e BGP presos em 58 s",
         "Era o tamanho da janela do teste. Investigar isso revelou a falha "
         "unilateral."),
        ("OSPF com a maior tabela",
         "Eu contava a RIB. Contando só as instaladas, RIP e OSPF empatam em "
         "15."),
    ]
    w, h, g = 5.3, 1.9, 0.27
    for i, (cab, corpo) in enumerate(itens):
        cartao(s, X0 + (i % 2) * (w + g), Y_CONTEUDO + (i // 2) * (h + g),
               w, h, cab, corpo, tam_cab=21, tam_corpo=17)
    rodape(s, "Dados da coleta com a janela curta preservados no repositório")
    notas(s, "Uns 40 segundos. Quatro coisas deram errado, e todas melhoraram o "
             "trabalho. RIP e OSPF não formavam adjacência no Docker: o "
             "netfilter da bridge descartava o multicast, e o BGP funcionava "
             "porque é TCP. A primeira medição deu o mesmo valor para os três "
             "protocolos, o que era impossível: eu media a primeira resposta "
             "depois do comando, e não a maior lacuna. Depois, RIP e BGP "
             "pareciam presos em 58 segundos, que era a janela do teste; "
             "investigar isso revelou a falha unilateral. E o OSPF parecia ter "
             "a maior tabela porque eu contava a RIB. Usei IA no trabalho, e "
             "cada hipótese dela foi verificada no laboratório; isso está "
             "documentado no repositório.")


def s_fim(prs, img):
    s = novo_slide(prs, FUNDO_FIM, img)
    texto(s, 6.53, 2.75, 6.2, 1.2, "Perguntas?", tamanho=66, negrito=True)
    texto(s, 6.53, 4.05, 5.9, 1.1,
          [[("./scripts/lab.sh up ospf", {"fonte": "Consolas", "tamanho": 18,
                                          "negrito": True})],
           [("Código, configurações e dados: [link do repositório no GitHub]",
             {"tamanho": 16, "cor": TEXTO_2})]], espaco_depois=6)
    notas(s, "Encerrar em 10 segundos e, se houver tempo, mostrar ao vivo: "
             "lab.sh up ospf, lab.sh vizinhos e o traceroute de h1 até "
             "10.3.0.10 indo pelo R5. Trocar o link do repositório antes da "
             "apresentação.")


# ===========================================================================

def main() -> None:
    d = json.loads(DADOS.read_text(encoding="utf-8"))
    with zipfile.ZipFile(MODELO) as z:
        imagens = {n: z.read(n) for n in (FUNDO_CAPA, FUNDO_CONTEUDO,
                                          FUNDO_DESTAQUE, FUNDO_FIM)}
    imagens[FUNDO_CONTEUDO] = limpar_fundo_conteudo(imagens[FUNDO_CONTEUDO])

    prs = Presentation(str(MODELO))
    n_modelo = len(prs.slides)

    s_capa(prs, imagens)
    s_construido(prs, imagens)
    s_plataforma(prs, imagens)
    s_topologia(prs, imagens)
    s_selecao(prs, imagens, d)
    s_convfria(prs, imagens, d)
    s_controle(prs, imagens, d)
    s_tabelas(prs, imagens, d)
    s_reconv_tempo(prs, imagens, d)
    s_reconv_cenarios(prs, imagens, d)
    s_comparativo(prs, imagens)
    s_licoes(prs, imagens)
    s_fim(prs, imagens)

    # remove os slides de exemplo do modelo; as partes sem referencia nao
    # sao gravadas pelo python-pptx
    lista = prs.slides._sldIdLst
    for sld in list(lista)[:n_modelo]:
        prs.part.drop_rel(sld.rId)
        lista.remove(sld)

    prs.save(str(SAIDA))
    print(f"{SAIDA.name}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
